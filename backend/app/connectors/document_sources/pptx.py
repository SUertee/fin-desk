"""PPTX parser: python-pptx, each slide → Section; shapes → text/table/image.

Block.page = slide_no (PPT physical-page contract; chunker does not cross it).
Two passes: collect shapes + cross-slide text repeat counting (filter template
headers like school name/logo), then build the Document.
"""

from __future__ import annotations

import re
from collections import Counter
from io import BytesIO

from app.connectors.document_sources._tables import make_table
from app.knowledge.document_models import Block, Document, Section

PPTX_HEADER_REPEAT = 3  # text repeated across >= this many slides → template header
TITLE_MAX_LEN = 30


def parse_pptx(content: bytes) -> Document:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(BytesIO(content))
    root = Section(level=0, title=None, synthetic=True)

    slides_data: list[tuple[int, str, list[tuple[str, object]]]] = []
    text_counter: Counter[str] = Counter()
    for slide_no, slide in enumerate(prs.slides, start=1):
        title = _slide_title(slide, slide_no)
        items: list[tuple[str, object]] = []
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    items.append(("image", shape.image.blob))
                except Exception:
                    pass
            elif shape.has_table:
                tbl = shape.table
                rows = [[cell.text.strip() for cell in row.cells] for row in tbl.rows]
                if len(rows) >= 2:
                    items.append(("table", rows))
            elif shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    items.append(("text", text))
                    text_counter[text] += 1
        slides_data.append((slide_no, title, items))

    header_set = {t for t, n in text_counter.items() if n >= PPTX_HEADER_REPEAT}

    for slide_no, title, items in slides_data:
        sec = Section(level=1, title=title, page=slide_no)
        root.children.append(sec)
        for kind, data in items:
            if kind == "image":
                sec.blocks.append(Block(type="image", image_bytes=data, page=slide_no))
            elif kind == "table":
                rows = data
                table = make_table(rows, f"pptx_table_{slide_no}_{id(data)}", slide_no)
                sec.blocks.append(Block(type="table", table=table, page=slide_no))
            else:
                if data in header_set:
                    continue
                if data == title:
                    continue
                sec.blocks.append(Block(type="text", content=data, page=slide_no))
    return Document(root=root)


def _slide_title(slide, slide_no: int) -> str:
    if slide.shapes.title is not None:
        t = slide.shapes.title.text_frame.text.strip()
        if t and _is_title_like(t):
            return t
    best_size = 0.0
    best_text = ""
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            para_text = ""
            para_size = 0.0
            for run in para.runs:
                para_text += run.text
                size = run.font.size
                if size and size.pt > para_size:
                    para_size = size.pt
            para_text = para_text.strip()
            if (
                para_size > best_size
                and para_text
                and len(para_text) <= TITLE_MAX_LEN
                and _is_title_like(para_text)
            ):
                best_size = para_size
                best_text = para_text
    return best_text or f"幻灯片{slide_no}"


def _is_title_like(text: str) -> bool:
    return bool(re.search(r"[一-鿿a-zA-Z]", text))
