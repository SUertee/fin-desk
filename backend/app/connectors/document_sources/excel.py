"""Excel parser: openpyxl, each sheet → Section(title=sheet name); table + images.

read_only=False (default) so drawings load (ws._images populated); image bytes
are read from the original zip (openpyxl closes the archive after load).
"""

from __future__ import annotations

import zipfile
from io import BytesIO

from app.connectors.document_sources._tables import make_table, normalize_table_header
from app.knowledge.document_models import Block, Document, Section


def parse_excel(content: bytes) -> Document:
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(content), data_only=True)
    zf = zipfile.ZipFile(BytesIO(content))
    root = Section(level=0, title=None, synthetic=True)
    for ws in wb.worksheets:
        rows = [
            [str(c.value) if c.value is not None else "" for c in row]
            for row in ws.iter_rows()
        ]
        rows = normalize_table_header(rows)
        has_table = len(rows) >= 2
        has_images = bool(ws._images)
        if not has_table and not has_images:
            continue
        blocks: list[Block] = []
        if has_table:
            blocks.append(Block(type="table", table=make_table(rows, f"sheet:{ws.title}", None)))
        for im in ws._images:
            try:
                img_bytes = zf.read(im.path.lstrip("/"))
                blocks.append(Block(type="image", image_bytes=img_bytes))
            except Exception:
                pass
        root.children.append(Section(level=1, title=ws.title, blocks=blocks))
    wb.close()
    zf.close()
    return Document(root=root)
